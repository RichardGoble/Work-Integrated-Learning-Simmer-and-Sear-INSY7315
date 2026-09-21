import os
import sys

# Try importing pptx, or run subprocess to install it
try:
    import pptx
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

def add_slide(title_text, content_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = title_text
    tf = slide.shapes.placeholders[1].text_frame
    for i, line in enumerate(content_lines):
        if i == 0:
            tf.text = line
        else:
            p = tf.add_paragraph()
            p.text = line.strip()
            if line.startswith("  "):
                p.level = 1

# 1. Title Slide
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Simmer & Sear"
slide.placeholders[1].text = "CoreFive Digital\nTiaan Walters | Richard Goble | Jared Nel | Ross Wassenaar | Heinrich Rossouw"

# 2. Problem & Solution
add_slide("Problem & Solution", [
    "The Problem:",
    "  • Manual wholesale quoting processes",
    "  • Fragmented brand catalogues for retail and wholesale",
    "The Solution:",
    "  • Unified ASP.NET Core e-commerce platform",
    "  • Dynamic bulk-order threshold routing",
    "  • Secure admin management portal"
])

# 3. Front-End Development (Design & UX)
add_slide("Front-End Development (Design & UX)", [
    "Visual Design:",
    "  • Frozen Design System (Partial views, CSS tokens)",
    "  • Pixel-perfect match to Figma prototypes",
    "Accessibility & UX:",
    "  • WCAG 2.2 Compliant (Semantic landmarks, 4.5:1 contrast)",
    "  • Keyboard-only navigation support",
    "  • Responsive breakpoints (320px width & 200% zoom)"
])

# 4. Back-End Development (Core Architecture)
add_slide("Back-End Development (Core Architecture)", [
    "Architecture Decision (ADR):",
    "  • Single ASP.NET Core Deployable (MVC + API)",
    "  • Eliminates CORS plumbing and cross-app token issues",
    "Design Patterns Implemented:",
    "  • Repository, Service Layer, DTO, Factory Method, Singleton",
    "Verification:",
    "  • ArchitectureTests (NetArchTest) proves 3-tier boundary in code"
])

# 5. Back-End Development (Database & Auth)
add_slide("Back-End Development (Database & Auth)", [
    "Database Architecture:",
    "  • Consolidated all data to MongoDB Atlas",
    "  • Documents for records, MongoDB GridFS for images",
    "  • Compound indexes, Text indexes, and $jsonSchema validators",
    "Custom Authentication:",
    "  • Replaced Firebase SSO with Application-managed identity",
    "  • PasswordHasher<T> with PBKDF2 and per-user salts"
])

# 6. Hosting
add_slide("Hosting Strategy", [
    "Infrastructure:",
    "  • Azure App Service (B1 Basic plan, Linux)",
    "  • MongoDB Atlas for database",
    "Environment Strategy:",
    "  • Single hosted environment (Production)",
    "  • Local Docker (mongo:7) for development",
    "  • Minimizes deployment risk, secrets sprawl, and costs"
])

# 7. GitHub & Version Control
add_slide("GitHub & Version Control", [
    "Branching Strategy (GitFlow):",
    "  • Protected 'main' (production) and 'develop' (buffer) branches",
    "  • Short-lived feature branches",
    "Collaboration Rules:",
    "  • Draft PRs on first commit for early CI visibility",
    "  • Strict 2-hour review SLAs",
    "  • No direct commits to main"
])

# 8. CI/CD Pipelines
add_slide("CI/CD Pipelines", [
    "Continuous Integration (GitHub Actions):",
    "  • Automated Build & Unit Tests",
    "  • Integration Tests against real mongo:7 container",
    "  • 'gitleaks' to prevent committed secrets",
    "Continuous Deployment:",
    "  • Automated deploy to Azure App Service",
    "  • Post-deploy Automated Smoke Test (/health/ready)",
    "  • Fails pipeline if health check does not return 200 OK"
])

# 9. Architecture & System Design
add_slide("Architecture & System Design", [
    "Data Flow:",
    "  • Browser ➔ Azure Web App (MVC)",
    "  • Azure Web App ➔ REST API (Internal routing)",
    "  • REST API ➔ MongoDB Atlas",
    "Security Boundary:",
    "  • Database is never exposed to the public internet",
    "  • API is the sole gatekeeper for all data access"
])

# 10. Live Technical Demonstration
add_slide("Live Technical Demonstration", [
    "1. Security & Audit Logging (RBAC):",
    "  • Sales Admin attempting Master Admin action",
    "  • Result: 403 Forbidden + Immutable Audit Log Entry",
    "2. Core Business Rule (Bulk Threshold):",
    "  • Add commercial quantity to cart",
    "  • Intercept standard checkout",
    "  • Generate PDF quote via GridFS & prompt sales meeting"
])

# 11. Requirements Alignment
add_slide("Requirements Alignment", [
    "Traceability & ADR Justification:",
    "  • NFR: 'Passwords must be hashed, never plaintext'",
    "  • Shift from Firebase met this requirement via PasswordHasher",
    "  • Requirement: 'Manage Wholesale & Retail'",
    "  • Met via dynamic cart threshold logic",
    "  • Requirement: 'Secure Admin Portal & Auditability'",
    "  • Met via custom API-level RBAC & MongoDB Audit Log"
])

# 12. Conclusion
add_slide("Conclusion", [
    "Summary:",
    "  • Evolved into a secure, 3-tier enterprise application",
    "  • Consolidated database (MongoDB) & Auth reduced integration risk",
    "  • High-performance, accessible, and responsive UI",
    "  • Delivered exact fidelity to client designs",
    "",
    "Thank You. Any Questions?"
])

prs.save('Simmer_and_Sear_Presentation.pptx')
print("Successfully created Simmer_and_Sear_Presentation.pptx")
